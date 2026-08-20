"""
Every supported format, as real bytes, through the real parser.

The suite tested tools thoroughly and formats hardly at all, so the file-shaped bugs
survived longest and were the worst: a workbook whose second sheet silently overwrote the
first, an export whose title row became the column names, a CSV that turned into a single
258,000-character chunk. None of them raised. Each produced a table that looked fine and
answered wrongly.

Fixtures are generated here rather than committed, so the tests carry their own shapes and
the repository carries no binary blobs.
"""

import io
import unittest
from unittest.mock import AsyncMock

import numpy as np
import pandas as pd

from app.core.config import get_settings
from app.services.ingest import IngestionService
from app.services.tools import _create_data_profile

settings = get_settings()


def ingestor():
    svc = IngestionService.__new__(IngestionService)
    svc.vector_store = AsyncMock()
    svc.session_store = AsyncMock()
    return svc


def workbook(sheets: dict, title_rows: int = 0) -> bytes:
    """An .xlsx, optionally with junk rows above the header as every export has."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            if title_rows:
                pd.DataFrame([[f"Export line {i}"] for i in range(title_rows)]).to_excel(
                    writer, sheet_name=name, index=False, header=False
                )
            frame.to_excel(writer, sheet_name=name, index=False, startrow=title_rows)
    return buffer.getvalue()


def document(paragraphs, table=None) -> bytes:
    import docx

    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        added = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                added.rows[r].cells[c].text = str(value)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def portable_document(pages) -> bytes:
    """A real PDF with a text layer, which is what pdfplumber needs to find anything."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    buffer = io.BytesIO()
    with PdfPages(buffer) as pdf:
        for lines in pages:
            figure = plt.figure(figsize=(8.27, 11.69))
            figure.text(0.08, 0.94, "\n".join(lines), va="top",
                        family="monospace", fontsize=10)
            pdf.savefig(figure)
            plt.close(figure)
    return buffer.getvalue()


def ruled_pdf() -> bytes:
    """
    A PDF whose table is drawn with cell borders, as a report tool exports it.

    Prose in a PDF yields text and no table, which is correct; pdfplumber finds a grid
    only where one is ruled. Both shapes are exercised, because "chart the numbers in
    this PDF" is a headline feature and neither shape had ever been tested end to end.
    """
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    figure, axes = plt.subplots(figsize=(8.27, 11.69))
    axes.axis("off")
    table = axes.table(
        cellText=[["North", "1,200.50", "2025-01-15"],
                  ["South", "2,300.75", "2025-02-20"]],
        colLabels=["region", "amount", "when"],
        loc="upper center",
    )
    table.scale(1, 2)
    buffer = io.BytesIO()
    figure.savefig(buffer, format="pdf")
    plt.close(figure)
    return buffer.getvalue()


class CommaSeparatedFiles(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.svc = ingestor()

    async def table(self, name, raw):
        chunks, tables = self.svc._parse_csv(name, raw, "doc")
        self.assertTrue(tables, f"{name} produced no table")
        frame = list(tables.values())[0]
        return chunks, await self.svc._universal_data_cleaning(frame, "t")

    async def test_a_plain_csv_round_trips(self):
        _, frame = await self.table(
            "sales.csv", b"region,revenue,units\nNorth,1000,5\nSouth,2000,10\n"
        )
        self.assertEqual(list(frame.columns), ["region", "revenue", "units"])
        self.assertEqual(frame["revenue"].sum(), 3000)

    async def test_a_semicolon_csv_round_trips(self):
        _, frame = await self.table(
            "euro.csv", b"region;revenue\nNorth;1000\nSouth;2000\n"
        )
        self.assertEqual(frame.shape, (2, 2))
        self.assertEqual(frame["revenue"].sum(), 3000)

    async def test_a_utf8_bom_does_not_corrupt_the_first_column(self):
        """Excel writes a BOM; read naively it becomes part of the first column name."""
        raw = "﻿region,revenue\nNorth,1000\n".encode("utf-8")
        _, frame = await self.table("bom.csv", raw)
        self.assertIn("region", [str(c).strip().lstrip("﻿") for c in frame.columns])

    async def test_accented_and_arabic_values_survive(self):
        raw = "region,note\nMünchen,café\nالرياض,ملاحظة\n".encode("utf-8")
        _, frame = await self.table("unicode.csv", raw)
        self.assertIn("München", frame["region"].tolist())
        self.assertIn("الرياض", frame["region"].tolist())

    async def test_currency_text_becomes_numeric(self):
        raw = b'month,amount\nJan,"1,200.50"\nFeb,"$2,300.75"\n'
        _, frame = await self.table("money.csv", raw)
        self.assertAlmostEqual(frame["amount"].sum(), 3501.25, places=2)

    async def test_a_large_file_is_indexed_as_a_preview(self):
        """A whole dataframe as one chunk took a later RAG query to 8m22s."""
        frame = pd.DataFrame({
            "id": range(6000),
            "amount": np.arange(6000, dtype=float),
        })
        chunks, table = await self.table("big.csv", frame.to_csv(index=False).encode())

        self.assertEqual(len(table), 6000, "every row still reaches the dataframe")
        biggest = max(len(c["content"]) for c in chunks)
        self.assertLessEqual(biggest, settings.CHUNK_SIZE * 1.5, "chunks stay bounded")
        self.assertLess(sum(len(c["content"]) for c in chunks), 20_000)


class Workbooks(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.svc = ingestor()

    async def test_each_sheet_becomes_its_own_table(self):
        raw = workbook({
            "Q1": pd.DataFrame({"site": ["A"], "tickets": [10]}),
            "Q2": pd.DataFrame({"site": ["B"], "tickets": [20]}),
            "Q3": pd.DataFrame({"site": ["C"], "tickets": [30]}),
        })
        _, tables = self.svc._parse_excel("ops.xlsx", raw, "doc")
        self.assertEqual(len(tables), 3, f"got {list(tables)}")

    async def test_a_header_below_title_rows_is_found(self):
        """
        Taken as-is the title becomes the column names — ['Export line 0', 'Unnamed: 1']
        — and nothing in the sheet can be addressed by the name it actually has.
        """
        raw = workbook(
            {"Tickets": pd.DataFrame({"month": ["2025-01"], "tickets": [42]})},
            title_rows=2,
        )
        _, tables = self.svc._parse_excel("ops.xlsx", raw, "doc")
        frame = list(tables.values())[0]
        self.assertEqual(list(frame.columns), ["month", "tickets"])
        self.assertNotIn("Unnamed: 1", frame.columns)

    async def test_a_normal_header_on_the_first_row_still_works(self):
        raw = workbook({"S": pd.DataFrame({"a": [1, 2], "b": [3, 4]})})
        _, tables = self.svc._parse_excel("plain.xlsx", raw, "doc")
        self.assertEqual(list(list(tables.values())[0].columns), ["a", "b"])

    async def test_a_totals_row_is_removed_rather_than_summed_twice(self):
        """
        This once asserted the *profile* still flagged a totals row after cleaning, which
        was the bug wearing a test's clothes: cleaning removes the real TOTAL, leaving
        10/20/30, and 30 is exactly half of 60 — so the detector flagged an ordinary row
        and the assertion passed for the wrong reason. What matters is the arithmetic.
        """
        frame = pd.DataFrame({"site": ["A", "B", "C"], "tickets": [10.0, 20.0, 30.0]})
        frame.loc[len(frame)] = ["TOTAL", 60.0]
        raw = workbook({"Tickets": frame})

        _, tables = self.svc._parse_excel("ops.xlsx", raw, "doc")
        cleaned = await self.svc._universal_data_cleaning(list(tables.values())[0], "t")

        self.assertEqual(len(cleaned), 3, "the TOTAL line is still loaded as data")
        self.assertEqual(cleaned["tickets"].sum(), 60.0)
        self.assertNotIn("TOTAL", cleaned["site"].astype(str).tolist())
        self.assertNotIn("Totals row", _create_data_profile(cleaned, "tickets"))


class WordDocuments(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.svc = ingestor()

    async def test_prose_and_a_table_both_come_out(self):
        raw = document(
            ["The registration process is described below. " * 8],
            [["Month", "Pay"], ["Jan", "1,200.50"], ["Feb", "2,300.75"]],
        )
        chunks, tables = self.svc._parse_docx("doc.docx", raw, "doc")

        self.assertTrue(any("registration" in c["content"] for c in chunks))
        self.assertTrue(tables)
        cleaned = await self.svc._universal_data_cleaning(list(tables.values())[0], "t")
        self.assertAlmostEqual(cleaned[cleaned.columns[1]].sum(), 3501.25, places=2)

    async def test_a_document_with_no_table_still_yields_text(self):
        raw = document(["A policy document with no tables at all. " * 10])
        chunks, tables = self.svc._parse_docx("policy.docx", raw, "doc")
        self.assertTrue(chunks)
        self.assertFalse(tables)

    async def test_arabic_prose_survives(self):
        raw = document(["نص عربي طويل يصف إجراءات التسجيل في الوزارة. " * 8])
        chunks, _ = self.svc._parse_docx("ar.docx", raw, "doc")
        self.assertTrue(any("التسجيل" in c["content"] for c in chunks))


class PortableDocuments(unittest.IsolatedAsyncioTestCase):
    """The PDF path had never been exercised with an actual PDF."""

    def setUp(self):
        self.svc = ingestor()

    async def test_text_is_extracted_from_every_page(self):
        raw = portable_document([
            ["ESCALATION PROCEDURE", "Tier 1 resolves within 4 working hours."],
            ["PENALTY TERMS", "Late delivery incurs 2% per week, capped at 10%."],
        ])
        chunks, _ = self.svc._parse_pdf("ops.pdf", raw, "doc")
        text = " ".join(c["content"] for c in chunks)

        self.assertIn("ESCALATION", text)
        self.assertIn("capped at 10%", text, "the second page must be read too")

    async def test_pages_are_attributed(self):
        raw = portable_document([["Page one text."], ["Page two text."]])
        chunks, _ = self.svc._parse_pdf("two.pdf", raw, "doc")
        self.assertTrue(chunks)
        self.assertTrue(all(c["source"] == "two.pdf" for c in chunks))
        self.assertTrue(all("page" in c["metadata"] for c in chunks))

    async def test_a_pdf_with_no_text_layer_is_refused(self):
        """A scan is an image; reporting a successful upload of nothing is worse."""
        blank = portable_document([[""]])
        with self.assertRaises(ValueError):
            await self.svc.ingest_file("scan.pdf", "scan.pdf", blank)


class PlainText(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.svc = ingestor()

    async def test_a_text_file_is_chunked(self):
        raw = "\n\n".join(f"Paragraph {i}. Policy text goes here." for i in range(40))
        chunks, tables = self.svc._parse_text("notes.txt", raw.encode(), "doc")
        self.assertGreater(len(chunks), 1)
        self.assertFalse(tables)
        for chunk in chunks:
            self.assertLessEqual(len(chunk["content"]), settings.CHUNK_SIZE * 1.5)


class UnsupportedInput(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.svc = ingestor()

    async def test_each_unreadable_shape_is_refused_not_accepted_empty(self):
        cases = {
            "archive.zip": b"PK\x03\x04not really a zip",
            "picture.png": b"\x89PNG\r\n\x1a\n\x00\x00",
            "empty.csv": b"",
            "broken.xlsx": b"PK\x03\x04corrupt workbook",
        }
        for name, raw in cases.items():
            with self.subTest(file=name):
                with self.assertRaises(ValueError):
                    await self.svc.ingest_file("alice", name, raw)


if __name__ == "__main__":
    unittest.main()


class TotalsRowsNeverReachTheAnswer(unittest.IsolatedAsyncioTestCase):
    """
    Stating it in the profile was not enough. Asked for a total, the agent read the
    warning and answered 7,208 anyway — exactly double — having decided the summary line
    was a site called "Unknown". A row that is provably the sum of the others is an
    artefact of the export rather than an observation, so it is removed at ingestion and
    the removal is logged.
    """

    def setUp(self):
        self.svc = ingestor()

    async def clean(self, frame):
        return await self.svc._universal_data_cleaning(frame, "t")

    async def test_the_totals_row_is_removed(self):
        frame = pd.DataFrame({
            "site": ["A", "B", "C", "TOTAL"],
            "tickets": [10.0, 20.0, 30.0, 60.0],
            "resolved": [8.0, 18.0, 24.0, 50.0],
        })
        cleaned = await self.clean(frame)
        self.assertEqual(len(cleaned), 3)
        self.assertEqual(cleaned["tickets"].sum(), 60)
        self.assertNotIn("TOTAL", cleaned["site"].tolist())

    async def test_a_label_in_any_language_is_caught(self):
        """The test is arithmetic, so it does not depend on the word "total"."""
        frame = pd.DataFrame({
            "site": ["A", "B", "C", "الإجمالي"],
            "tickets": [10.0, 20.0, 30.0, 60.0],
            "resolved": [8.0, 18.0, 24.0, 50.0],
        })
        self.assertEqual(len(await self.clean(frame)), 3)

    async def test_ordinary_data_is_never_dropped(self):
        frame = pd.DataFrame({
            "site": ["A", "B", "C", "D"],
            "tickets": [10.0, 20.0, 30.0, 40.0],
            "resolved": [8.0, 18.0, 24.0, 33.0],
        })
        self.assertEqual(len(await self.clean(frame)), 4)

    async def test_one_column_agreeing_by_chance_is_not_enough(self):
        """In [1, 2, 3, 6] the six equals the sum of the rest, and is real data."""
        frame = pd.DataFrame({
            "site": ["A", "B", "C", "D"],
            "tickets": [1.0, 2.0, 3.0, 6.0],
            "resolved": [5.0, 9.0, 2.0, 7.0],
        })
        self.assertEqual(len(await self.clean(frame)), 4)

    async def test_an_exported_workbook_totals_correctly_end_to_end(self):
        rows = pd.DataFrame({
            "month": [f"2025-{m:02d}" for m in range(1, 4)],
            "site": ["Riyadh", "Jeddah", "Dammam"],
            "tickets": [100.0, 200.0, 300.0],
            "resolved": [90.0, 180.0, 270.0],
        })
        rows.loc[len(rows)] = ["TOTAL", "", 600.0, 540.0]
        raw = workbook({"Tickets": rows}, title_rows=2)

        _, tables = self.svc._parse_excel("ops.xlsx", raw, "doc")
        cleaned = await self.clean(list(tables.values())[0])

        self.assertEqual(list(cleaned.columns), ["month", "site", "tickets", "resolved"])
        self.assertEqual(cleaned["tickets"].sum(), 600, "not 1,200")


class EveryFormatClearsEveryStage(unittest.IsolatedAsyncioTestCase):
    """
    The same questions asked of all six formats, rather than of the convenient one.

    Parsing was tested per format and the rest of the pipeline only through CSV, so
    "which formats have actually been through cleaning, profiling, indexing, a tool and
    a delete" had no answer. Tabulating it is how the gaps become visible at all.
    """

    TABLE_BEARING = {"csv", "xlsx", "xls", "docx", "pdf"}

    def files(self):
        rows = [{"region": "North", "amount": "1,200.50", "when": "2025-01-15"},
                {"region": " South ", "amount": "$2,300.75", "when": "2025-02-20"}]
        return {
            "csv": ("data.csv", pd.DataFrame(rows).to_csv(index=False).encode()),
            "xlsx": ("data.xlsx", workbook({"Data": pd.DataFrame(rows)})),
            "xls": ("data.xls", workbook({"Data": pd.DataFrame(rows)})),
            "txt": ("notes.txt", ("Quarterly review. " * 40).encode()),
            "docx": ("report.docx", document(
                ["Quarterly review of regional performance. " * 8],
                [["region", "amount", "when"],
                 ["North", "1,200.50", "2025-01-15"],
                 ["South", "2,300.75", "2025-02-20"]],
            )),
            "pdf": ("report.pdf", ruled_pdf()),
        }

    async def pipeline(self, name, data):
        from app.services.ingest import IngestionService
        from app.services.tools import _create_data_profile, execute_tool
        from tests.test_stores import session_store, vector_store

        sessions, vectors = session_store(), vector_store()
        svc = IngestionService.__new__(IngestionService)
        svc.session_store, svc.vector_store = sessions, vectors

        out = await svc.ingest_file("s", name, data)
        stages = {
            "parse": out["status"] == "success",
            "chunk": out["text_chunks"] > 0,
            "index": vectors.index.ntotal == out["text_chunks"],
        }

        names = await sessions.list_dataframes("s")
        stages["table"] = bool(names)
        if names:
            frame = await sessions.get_dataframe("s", names[0])
            amount = next(c for c in frame.columns if "amount" in str(c).lower())
            stages["clean"] = (
                pd.api.types.is_numeric_dtype(frame[amount])
                and abs(frame[amount].sum() - 3501.25) < 0.01
            )
            stages["profile"] = "Column Types" in _create_data_profile(frame, names[0])
            said = await execute_tool(
                "calculate_statistics",
                {"table_name": names[0], "operation": "sum", "column": amount},
                {names[0]: frame},
            )
            stages["tool"] = "3501.25" in said.replace(",", "")

        docs = await sessions.get_session_documents("s")
        await sessions.delete_document("s", docs[0]["doc_id"], vectors)
        stages["delete"] = (
            vectors.index.ntotal == 0 and await sessions.list_dataframes("s") == []
        )
        return stages

    async def test_every_format_parses_chunks_indexes_and_deletes(self):
        for ext, (name, data) in self.files().items():
            with self.subTest(format=ext):
                stages = await self.pipeline(name, data)
                for stage in ("parse", "chunk", "index", "delete"):
                    self.assertTrue(stages[stage], f"{ext} failed at {stage}")

    async def test_every_tabular_format_reaches_a_tool_with_clean_numbers(self):
        """Currency written as text must survive the whole way to an aggregate."""
        for ext, (name, data) in self.files().items():
            if ext not in self.TABLE_BEARING:
                continue
            with self.subTest(format=ext):
                stages = await self.pipeline(name, data)
                for stage in ("table", "clean", "profile", "tool"):
                    self.assertTrue(stages[stage], f"{ext} failed at {stage}")

    async def test_a_prose_only_format_yields_no_table_rather_than_an_empty_one(self):
        stages = await self.pipeline(*self.files()["txt"])
        self.assertTrue(stages["chunk"])
        self.assertFalse(stages["table"])


class EveryEncodingAndDelimiterAFileMightUse(unittest.IsolatedAsyncioTestCase):
    """
    Encoding was not attempted at all: `pd.read_csv` was handed raw bytes and whatever it
    guessed was final. Excel's "Unicode Text" export is UTF-16 and a European Excel writes
    cp1252 — both raised UnicodeDecodeError, and because a parser that extracts nothing
    is treated as a corrupt file, the upload was refused as unreadable. The user's file was
    fine; the reader was not.

    UTF-16 is tried only behind its byte-order mark. Without one it is undetectable — it
    decodes almost any even-length input into mojibake, and doing so swallowed a valid
    cp1252 file and produced an empty table.
    """

    def setUp(self):
        self.svc = ingestor()

    async def frame_from(self, raw):
        _, tables = self.svc._parse_csv("f.csv", raw, "doc")
        self.assertTrue(tables, "nothing parsed")
        return list(tables.values())[0]

    async def test_utf16_with_a_bom(self):
        frame = await self.frame_from("site,value\nA,10\nB,20\n".encode("utf-16"))
        self.assertEqual(frame.shape, (2, 2))
        self.assertEqual(list(frame.columns), ["site", "value"])

    async def test_cp1252_accents(self):
        frame = await self.frame_from(
            "site,value\nCafé,10\nMünchen,20\n".encode("cp1252"))
        self.assertEqual(frame.shape, (2, 2))
        self.assertIn("Café", frame["site"].tolist())

    async def test_a_utf8_bom_does_not_leak_into_the_first_column(self):
        frame = await self.frame_from("﻿site,value\nA,10\n".encode("utf-8"))
        self.assertEqual(list(frame.columns), ["site", "value"])

    async def test_every_common_delimiter_splits(self):
        for sep in (",", ";", "\t", "|"):
            with self.subTest(sep=sep):
                raw = f"site{sep}value\nA{sep}10\nB{sep}20\n".encode()
                self.assertEqual((await self.frame_from(raw)).shape, (2, 2))

    async def test_a_delimiter_inside_a_quoted_value_is_not_a_split(self):
        frame = await self.frame_from(
            b'site,value\n"Riyadh, KSA",10\n"Jeddah, KSA",20\n')
        self.assertEqual(frame.shape, (2, 2))
        self.assertIn("Riyadh, KSA", frame["site"].tolist())

    async def test_crlf_line_endings(self):
        frame = await self.frame_from(b"site,value\r\nA,10\r\nB,20\r\n")
        self.assertEqual(frame.shape, (2, 2))


class ColumnNamesThatAreNotStrings(unittest.IsolatedAsyncioTestCase):
    """
    A header row of years loads as integers, and date detection called `col.lower()` on
    one. AttributeError took down the whole file — every column, not the one it tripped on.
    """

    async def test_integer_column_names_do_not_crash_cleaning(self):
        svc = ingestor()
        frame = pd.DataFrame({2023: ["a", "b"], 2024: [10, 20]})
        cleaned = await svc._universal_data_cleaning(frame, "t")
        self.assertEqual(len(cleaned), 2)

    async def test_a_column_named_with_a_float_survives(self):
        svc = ingestor()
        frame = pd.DataFrame({1.5: ["a", "b"], "v": [10, 20]})
        cleaned = await svc._universal_data_cleaning(frame, "t")
        self.assertEqual(len(cleaned), 2)
