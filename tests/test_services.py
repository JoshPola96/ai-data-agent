"""
The service layer under regression: ingestion, retrieval, response parsing, storage.

These cover the paths where a fault is invisible from the outside — a sheet that never
arrived, a chunk with the wrong owner, a cached index answering for a corpus that has
changed. Each one produced a plausible answer while being wrong.
"""

import asyncio
import io
import json
import pickle
import unittest
from unittest.mock import AsyncMock, patch

import numpy as np
import pandas as pd

import app.main as main
import app.services.retrieval as retr_mod
from app.core.config import get_settings
from app.core.session import SessionStore
from app.services.ingest import IngestionService
from app.services.llm import _is_transient, _provider_reason
from app.services.retrieval import HybridRetriever

settings = get_settings()


def ingestor():
    """An IngestionService with its two stores mocked; only parsing is under test."""
    svc = IngestionService.__new__(IngestionService)
    svc.vector_store = AsyncMock()
    svc.session_store = AsyncMock()
    return svc


class ChunkingPreservesTheDocument(unittest.TestCase):
    def setUp(self):
        self.svc = ingestor()
        body = "Sentence one is here. Sentence two follows on. Sentence three closes."
        self.text = "\n\n".join(f"Paragraph {i}. {body}" for i in range(40))
        self.chunks = self.svc._semantic_chunking(self.text, "t.txt", "d1", 1)

    def test_a_long_document_splits(self):
        self.assertGreater(len(self.chunks), 1)

    def test_no_chunk_wildly_exceeds_the_configured_size(self):
        biggest = max(len(c["content"]) for c in self.chunks)
        self.assertLess(biggest, settings.CHUNK_SIZE * 1.5)

    def test_no_paragraph_is_lost_between_chunks(self):
        """A paragraph that falls in a gap is unreachable by retrieval, silently."""
        joined = " ".join(c["content"] for c in self.chunks)
        missing = [i for i in range(40) if f"Paragraph {i}." not in joined]
        self.assertEqual(missing, [])

    def test_consecutive_chunks_overlap(self):
        for first, second in zip(self.chunks, self.chunks[1:]):
            tail = first["content"][-settings.CHUNK_OVERLAP:]
            words = [w for w in tail.split() if len(w) > 4]
            head = second["content"][: settings.CHUNK_OVERLAP + 80]
            self.assertTrue(any(w in head for w in words))

    def test_an_unbroken_paragraph_is_not_truncated(self):
        big = "x" * (settings.CHUNK_SIZE * 3)
        chunks = self.svc._semantic_chunking(big, "t", "d", 1)
        kept = sum(len(c["content"]) for c in chunks)
        self.assertGreaterEqual(kept, len(big) * 0.9)

    def test_a_document_shorter_than_the_minimum_is_still_kept(self):
        chunks = self.svc._semantic_chunking("Short.", "t", "d", 1)
        self.assertEqual([c["content"] for c in chunks], ["Short."])

    def test_empty_input_yields_nothing(self):
        self.assertEqual(self.svc._semantic_chunking("", "t", "d", 1), [])
        self.assertEqual(self.svc._semantic_chunking("  \n\n ", "t", "d", 1), [])

    def test_arabic_chunks_without_mangling(self):
        arabic = "\n\n".join(["الفقرة الأولى هنا. نص عربي طويل يتكرر في المستند." * 3] * 12)
        chunks = self.svc._semantic_chunking(arabic, "ar.txt", "d", 1)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(c["content"].strip() for c in chunks))

    def test_attribution_travels_with_every_chunk(self):
        """Retrieval scoping and citations both read these fields."""
        for chunk in self.svc._semantic_chunking(self.text, "report.pdf", "doc-9", 7):
            self.assertEqual(chunk["source"], "report.pdf")
            self.assertEqual(chunk["doc_id"], "doc-9")
            self.assertEqual(chunk["metadata"]["page"], 7)
            self.assertEqual(chunk["metadata"]["chunk_size"], len(chunk["content"]))


class ParsersExtractWhatIsThere(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.svc = ingestor()

    def table(self, result):
        _, tables = result
        self.assertTrue(tables, "expected at least one table")
        return list(tables.values())[0]

    async def test_a_csv_becomes_a_usable_table(self):
        raw = b"region,revenue,units\nNorth,1000,5\nSouth,2000,10\n"
        df = self.table(self.svc._parse_csv("sales.csv", raw, "d1"))
        self.assertEqual(list(df.columns), ["region", "revenue", "units"])
        self.assertEqual(len(df), 2)
        clean = await self.svc._universal_data_cleaning(df, "sales")
        self.assertEqual(clean["revenue"].sum(), 3000)

    async def test_a_semicolon_csv_is_not_read_as_one_column(self):
        """Excel writes semicolons across most of Europe; one column is unanalysable."""
        raw = b"region;revenue\nNorth;1000\nSouth;2000\n"
        df = self.table(self.svc._parse_csv("euro.csv", raw, "d2"))
        self.assertEqual(df.shape[1], 2, f"columns: {list(df.columns)}")

    async def test_blank_lines_do_not_become_rows(self):
        raw = b"region,revenue\n\nNorth,1000\n\n\nSouth,2000\n"
        self.assertEqual(len(self.table(self.svc._parse_csv("gaps.csv", raw, "d3"))), 2)

    async def test_every_worksheet_becomes_its_own_table(self):
        """
        Table names were built by cleaning "book.xlsx_Sheet", and cleaning split on the
        first dot — so every sheet collapsed to "book" and all but the last were
        overwritten. A two-sheet workbook silently lost half its data.
        """
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            pd.DataFrame({"a": [1, 2]}).to_excel(writer, sheet_name="First", index=False)
            pd.DataFrame({"b": [3]}).to_excel(writer, sheet_name="Second", index=False)
        _, tables = self.svc._parse_excel("book.xlsx", buf.getvalue(), "d4")
        self.assertEqual(len(tables), 2, f"tables: {list(tables)}")
        self.assertTrue(all(not t.empty for t in tables.values()))

    async def test_a_docx_yields_prose_and_tables(self):
        import docx

        document = docx.Document()
        document.add_paragraph("This document explains the registration process. " * 6)
        table = document.add_table(rows=3, cols=2)
        for r, row in enumerate([["Month", "Pay"], ["Jan", "1,200.50"], ["Feb", "2,300.75"]]):
            for c, value in enumerate(row):
                table.rows[r].cells[c].text = value
        buf = io.BytesIO()
        document.save(buf)

        chunks, tables = self.svc._parse_docx("doc.docx", buf.getvalue(), "d5")
        self.assertTrue(any("registration" in c["content"] for c in chunks))
        self.assertTrue(tables)

        extracted = list(tables.values())[0]
        self.assertEqual(list(extracted.columns)[:2], ["Month", "Pay"])
        clean = await self.svc._universal_data_cleaning(extracted, "pay")
        self.assertAlmostEqual(clean[clean.columns[1]].sum(), 3501.25, places=2)

    async def test_an_unreadable_file_fails_instead_of_reporting_success(self):
        """
        Parsers return empty on failure, so a corrupt upload used to be reported as a
        success with nothing in it — the file sat in the sidebar looking ingested.
        """
        with self.assertRaises(ValueError) as caught:
            await self.svc.ingest_file("alice", "thing.zip", b"PK\x03\x04garbage")
        self.assertIn("Nothing could be extracted", str(caught.exception))


class NamesStayAddressable(unittest.TestCase):
    def setUp(self):
        self.svc = ingestor()

    def test_only_a_trailing_extension_is_stripped(self):
        self.assertEqual(self.svc._clean_name("2025.01.15 report.csv"), "2025_01_15_report")

    def test_punctuation_becomes_underscores(self):
        self.assertEqual(self.svc._clean_name("My Report (final).xlsx"), "My_Report__final")

    def test_a_name_with_no_ascii_still_produces_one(self):
        """Sanitising an Arabic filename left the empty string — an unaddressable table."""
        name = self.svc._clean_name("تقرير.docx")
        self.assertTrue(name)
        self.assertEqual(name, self.svc._clean_name("تقرير.docx"), "must be stable")

    def test_doc_ids_are_unique_per_upload(self):
        ids = {self.svc._generate_doc_id("same.csv") for _ in range(50)}
        self.assertEqual(len(ids), 50)
        self.assertTrue(all(i.startswith("same_") for i in ids))


class OwnershipIsStampedOnIngest(unittest.IsolatedAsyncioTestCase):
    async def test_each_upload_carries_only_its_own_session(self):
        """This stamp is the whole of retrieval scoping and session deletion."""
        svc = ingestor()
        await svc.ingest_file("alice", "a.csv", b"region,revenue\nNorth,1\n")
        await svc.ingest_file("bob", "b.txt", b"Prose about revenue in the north. " * 30)

        batches = [c.args[0] for c in svc.vector_store.add_documents.call_args_list]
        self.assertEqual([{c["session_id"] for c in b} for b in batches],
                         [{"alice"}, {"bob"}])

    async def test_the_registry_records_the_tables_a_file_produced(self):
        svc = ingestor()
        await svc.ingest_file("alice", "sales.csv", b"region,revenue\nNorth,1\n")
        metadata = svc.session_store.register_document.call_args.args[3]
        self.assertEqual(metadata["tables"], ["sales"])


DOCS = [
    {"content": "alice quarterly revenue report northern region", "session_id": "alice", "source": "a1.pdf"},
    {"content": "alice invoice INV-88421 paid in full", "session_id": "alice", "source": "a2.pdf"},
    {"content": "bob confidential salary schedule senior staff", "session_id": "bob", "source": "b1.pdf"},
    {"content": "bob expense claim travel southern branch", "session_id": "bob", "source": "b2.pdf"},
    {"content": "shared employee handbook leave and conduct", "session_id": "kb_default", "source": "kb.pdf"},
]


class FakeStore:
    """Stands in for the vector store: semantic search answers with global positions."""

    def __init__(self, docs):
        self.docs = list(docs)
        self.version = 1
        self.search_calls = []

    async def get_documents(self):
        return self.docs

    async def search(self, query, k, allowed=None):
        self.search_calls.append({"query": query, "k": k, "allowed": allowed})
        terms = set(query.lower().split())
        scored = [
            (i, float(len(terms & set(d["content"].lower().split()))))
            for i, d in enumerate(self.docs)
            if allowed is None or i in allowed
        ]
        scored.sort(key=lambda pair: -pair[1])
        return scored[:k]


class RetrievalStaysInItsLane(unittest.IsolatedAsyncioTestCase):
    def retriever(self, store, reranker=None):
        r = HybridRetriever.__new__(HybridRetriever)
        r.vector_store = store
        r._bm25 = None
        r._bm25_version = None
        r._bm25_lock = asyncio.Lock()
        r.reranker = reranker
        r._reranker_initialized = reranker is not None
        return r

    async def retrieve(self, retriever, *args, rerank=False, threshold=0.5, **kwargs):
        with patch.object(retr_mod.settings, "USE_RERANKING", rerank), \
             patch.object(retr_mod.settings, "RERANK_THRESHOLD", threshold):
            return await retriever.retrieve(*args, **kwargs)

    async def test_another_sessions_documents_are_never_returned(self):
        store = FakeStore(DOCS)
        hits = await self.retrieve(
            self.retriever(store), "confidential salary",
            session_ids={"alice", "kb_default"},
        )
        self.assertNotIn("bob", {h["session_id"] for h in hits})
        self.assertFalse(any("salary schedule" in h["content"] for h in hits))

    async def test_a_session_can_reach_its_own_documents(self):
        hits = await self.retrieve(
            self.retriever(FakeStore(DOCS)), "confidential salary", session_ids={"bob"}
        )
        self.assertTrue(hits)
        self.assertEqual({h["session_id"] for h in hits}, {"bob"})

    async def test_a_session_with_nothing_retrieves_nothing(self):
        hits = await self.retrieve(
            self.retriever(FakeStore(DOCS)), "revenue", session_ids={"stranger"}
        )
        self.assertEqual(hits, [])

    async def test_scoping_is_pushed_into_the_search_not_applied_after(self):
        """Filtering after the fact would return fewer than top_k on a shared corpus."""
        store = FakeStore(DOCS)
        await self.retrieve(self.retriever(store), "revenue", session_ids={"alice"})
        call = store.search_calls[-1]
        self.assertEqual(call["allowed"], {0, 1})
        self.assertGreaterEqual(call["k"], settings.RETRIEVAL_CANDIDATES)

    async def test_a_hit_keeps_the_source_of_its_own_document(self):
        """An off-by-one in position mapping would cite the neighbouring document."""
        hits = await self.retrieve(
            self.retriever(FakeStore(DOCS)), "INV-88421", session_ids={"alice"}
        )
        self.assertIn("INV-88421", hits[0]["content"])
        by_content = {d["content"]: d["source"] for d in DOCS}
        for hit in hits:
            self.assertEqual(hit["source"], by_content[hit["content"]])

    async def test_an_unchanged_corpus_reuses_the_cached_index(self):
        store = FakeStore(DOCS)
        r = self.retriever(store)
        await self.retrieve(r, "handbook", session_ids=None)
        first = r._bm25
        await self.retrieve(r, "handbook", session_ids=None)
        self.assertIs(r._bm25, first)

    async def test_a_new_document_is_findable_immediately(self):
        """A cache keyed on nothing would answer for a corpus that no longer exists."""
        store = FakeStore(DOCS)
        r = self.retriever(store)
        await self.retrieve(r, "handbook", session_ids=None)
        stale = r._bm25

        store.docs = DOCS + [{"content": "newly added quarterly forecast eastern region",
                              "session_id": "alice", "source": "new.pdf"}]
        store.version = 2
        hits = await self.retrieve(r, "forecast eastern", session_ids={"alice"})

        self.assertIsNot(r._bm25, stale)
        self.assertTrue(any("forecast" in h["content"] for h in hits))

    async def test_after_a_removal_only_survivors_are_returned(self):
        store = FakeStore(DOCS)
        r = self.retriever(store)
        await self.retrieve(r, "handbook", session_ids=None)
        store.docs = DOCS[:2]
        store.version = 5
        hits = await self.retrieve(r, "handbook", session_ids=None)
        surviving = {d["content"] for d in store.docs}
        self.assertTrue(all(h["content"] in surviving for h in hits))

    async def test_no_more_than_top_k_is_returned(self):
        hits = await self.retrieve(self.retriever(FakeStore(DOCS)), "the", session_ids=None)
        self.assertLessEqual(len(hits), settings.RETRIEVAL_TOP_K)

    async def test_an_empty_store_retrieves_nothing(self):
        self.assertEqual(
            await self.retrieve(self.retriever(FakeStore([])), "anything", session_ids=None),
            [],
        )

    async def test_a_document_matching_several_variants_appears_once(self):
        store = FakeStore(DOCS)
        hits = await self.retrieve(
            self.retriever(store), "revenue",
            expanded_queries=["revenue", "quarterly income", "revenue"],
            session_ids={"alice"},
        )
        contents = [h["content"] for h in hits]
        self.assertEqual(len(contents), len(set(contents)))
        self.assertEqual(len(store.search_calls), 3)

    async def test_documents_below_the_rerank_threshold_are_dropped(self):
        class Reranker:
            def predict(self, pairs, show_progress_bar=False):
                return np.array([10.0, 9.0, -5.0, -6.0, -7.0][: len(pairs)])

        hits = await self.retrieve(
            self.retriever(FakeStore(DOCS), reranker=Reranker()),
            "handbook conduct", session_ids=None, rerank=True, threshold=0.5,
        )
        self.assertTrue(hits)
        self.assertTrue(all(h.get("rerank_score", 1) >= 0.5 for h in hits))


class FusionRanksSensibly(unittest.TestCase):
    def setUp(self):
        self.r = HybridRetriever.__new__(HybridRetriever)

    def test_every_candidate_survives_fusion(self):
        self.assertEqual(set(self.r._rrf_fusion([(7, 9.9), (3, 5.0)], [(3, 0.9), (7, 0.4)])),
                         {3, 7})

    def test_the_heavier_weight_breaks_a_tie(self):
        order = self.r._rrf_fusion([(7, 9.9), (3, 5.0)], [(3, 0.9), (7, 0.4)])
        winner = 3 if settings.SEMANTIC_WEIGHT >= settings.BM25_WEIGHT else 7
        self.assertEqual(order[0], winner)

    def test_agreement_between_both_searches_ranks_first(self):
        self.assertEqual(self.r._rrf_fusion([(4, 1.0), (5, 0.5)], [(4, 1.0), (5, 0.5)])[0], 4)

    def test_one_sided_results_keep_their_order(self):
        self.assertEqual(self.r._rrf_fusion([(1, 1.0), (2, 0.5)], []), [1, 2])

    def test_nothing_fuses_to_nothing(self):
        self.assertEqual(self.r._rrf_fusion([], []), [])


GOOD_JSON = {
    "answer": "East leads on revenue.",
    "visualizations": [],
    "key_insights": ["East is highest"],
    "sources_used": ["sales"],
}


class ResponseParsingDegradesGracefully(unittest.TestCase):
    def extract(self, content):
        return main.extract_structured_response(content)

    def test_plain_json_is_parsed(self):
        text, meta = self.extract(json.dumps(GOOD_JSON))
        self.assertEqual(text, GOOD_JSON["answer"])
        self.assertEqual(meta["key_insights"], GOOD_JSON["key_insights"])
        self.assertEqual(meta["sources_from_response"], ["sales"])

    def test_fenced_blocks_are_parsed(self):
        for fence in ("```json\n{body}\n```", "```\n{body}\n```",
                      "Here you go:\n```json\n{body}\n```\nHope that helps."):
            with self.subTest(fence=fence[:12]):
                content = fence.format(body=json.dumps(GOOD_JSON))
                self.assertEqual(self.extract(content)[0], GOOD_JSON["answer"])

    def test_a_double_encoded_payload_is_unwrapped(self):
        self.assertEqual(self.extract(json.dumps(json.dumps(GOOD_JSON)))[0],
                         GOOD_JSON["answer"])

    def test_non_latin_and_emoji_survive(self):
        for answer in ("الإجابة هنا بالعربية", "Revenue rose 📈 by 12%", "The region's total"):
            with self.subTest(answer=answer[:12]):
                payload = json.dumps(dict(GOOD_JSON, answer=answer), ensure_ascii=False)
                self.assertEqual(self.extract(payload)[0], answer)

    def test_unstructured_prose_still_reaches_the_user(self):
        text, meta = self.extract("The East region leads with 317,000 in revenue.")
        self.assertIn("East", text)
        self.assertEqual(meta["key_insights"], [])
        self.assertEqual(meta["visualizations"], [])

    def test_broken_payloads_do_not_raise(self):
        for content in ("", '{"answer": "half', json.dumps({"key_insights": ["no answer"]})):
            with self.subTest(content=content[:14]):
                text, meta = self.extract(content)
                self.assertIsInstance(text, str)
                self.assertIsInstance(meta.get("visualizations"), list)


class RetryClassification(unittest.TestCase):
    class Fault(Exception):
        def __init__(self, status=None, code=None):
            super().__init__("provider said no")
            if status is not None:
                self.status_code = status
            if code is not None:
                self.code = code

    def test_transient_statuses_are_retried(self):
        for status in (408, 429, 500, 502, 503, 504):
            with self.subTest(status=status):
                self.assertTrue(_is_transient(self.Fault(status=status)))

    def test_permanent_statuses_are_not(self):
        for status in (400, 401, 403, 404, 422):
            with self.subTest(status=status):
                self.assertFalse(_is_transient(self.Fault(status=status)))

    def test_the_google_style_code_attribute_is_read(self):
        self.assertTrue(_is_transient(self.Fault(code=429)))

    def test_a_status_free_exception_is_not_retried(self):
        self.assertFalse(_is_transient(Exception("mysterious")))

    def test_a_string_status_is_not_mistaken_for_a_code(self):
        self.assertFalse(_is_transient(self.Fault(status="429")))

    def test_the_reason_shown_to_the_user_is_short_and_specific(self):
        self.assertIn("429", _provider_reason(self.Fault(status=429)))
        self.assertLess(len(_provider_reason(Exception("x" * 400))), 200)
        self.assertTrue(_provider_reason(None))


class StorageKeepsItsPromises(unittest.TestCase):
    def test_a_dataframe_survives_the_round_trip_unchanged(self):
        """Losing a dtype in transit turns every later numeric answer into a string one."""
        df = pd.DataFrame({
            "region": ["North", "South"],
            "revenue": [1000.5, 2000.25],
            "units": np.array([5, 10], dtype="int64"),
            "when": pd.to_datetime(["2025-01-15", "2025-02-15"]),
            "flag": [True, False],
            "missing": [np.nan, 3.0],
            "arabic": ["شمال", "جنوب"],
        })
        revived = pickle.loads(pickle.dumps(df))
        self.assertTrue(revived.equals(df))
        self.assertEqual(list(revived.dtypes), list(df.dtypes))
        self.assertEqual(revived["revenue"].sum(), 3000.75)
        self.assertEqual(revived["arabic"].tolist(), ["شمال", "جنوب"])

    def test_two_sessions_cannot_collide_on_a_table_name(self):
        store = SessionStore.__new__(SessionStore)
        alice, bob = store._df_key("alice", "sales"), store._df_key("bob", "sales")
        self.assertNotEqual(alice, bob)
        self.assertIn("alice", alice)
        self.assertNotEqual(store._df_key("alice", "sales:evil"), alice)


if __name__ == "__main__":
    unittest.main()


class FailoverCarriesTheConversation(unittest.IsolatedAsyncioTestCase):
    """
    A thought_signature belongs to the model that issued it.

    Gemini rejects a replayed function call whose signature is missing, and equally
    rejects one carrying another model's — so a conversation holding tool calls could not
    be handed over. Observed live: the primary timed out on turn 2, the fallback was sent
    turn 1's function call without a signature, and the request died on
    "400 Function call is missing a thought_signature in functionCall parts" with no
    fallback left. The answer the user got was an error.
    """

    HISTORY = [
        {"role": "user", "content": "which categories are unprofitable?"},
        {"role": "assistant", "content": "", "tool_calls": [{
            "id": "c1", "name": "query_data", "arguments": {"table_name": "sales"},
            "thought_signature": b"sig-from-2.5",
            "signature_model": "gemini:gemini-2.5-flash",
        }]},
        {"role": "tool", "name": "query_data", "content": "Query Results (3 rows)"},
    ]

    async def parts_sent_to(self, model):
        from unittest.mock import MagicMock

        from app.services.llm import LLMService

        service = LLMService()
        service._initialized = True
        calls = []

        answer = json.dumps({"answer": "ok", "visualizations": [],
                             "key_insights": [], "sources_used": []})

        class Models:
            def generate_content(self, model, contents, config):
                calls.append(contents)
                part = MagicMock(text=answer, function_call=None, thought_signature=None)
                candidate = MagicMock()
                candidate.content.parts = [part]
                return MagicMock(candidates=[candidate])

        service.gemini_client = MagicMock()
        service.gemini_client.models = Models()

        await service._chat_gemini(self.HISTORY, "sys", None, 0.1, True, model)
        return [part for content in calls[0] for part in content.parts]

    async def test_the_issuing_model_replays_calls_natively(self):
        parts = await self.parts_sent_to("gemini:gemini-2.5-flash")
        calls = [p for p in parts if getattr(p, "function_call", None)]
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0].thought_signature, b"sig-from-2.5")
        self.assertTrue([p for p in parts if getattr(p, "function_response", None)])

    async def test_another_model_receives_the_exchange_as_text(self):
        parts = await self.parts_sent_to("gemini:gemini-flash-latest")
        self.assertEqual([p for p in parts if getattr(p, "function_call", None)], [])
        self.assertEqual([p for p in parts if getattr(p, "function_response", None)], [])

        text = " ".join(p.text for p in parts if getattr(p, "text", None))
        self.assertIn("query_data", text, "the call itself must survive the handover")
        self.assertIn("Query Results (3 rows)", text, "so must what it returned")

    async def test_the_issuing_model_replays_natively_without_a_signature(self):
        """
        gemini-2.5-flash sends no thought_signature and needs none. Treating "no
        signature" as "not ours" degraded every multi-turn request on that model to text,
        losing the distinction between a tool result and a paragraph about one.
        """
        history = [
            {"role": "user", "content": "totals please"},
            {"role": "assistant", "content": "", "tool_calls": [{
                "id": "c1", "name": "query_data", "arguments": {"table_name": "sales"},
                "thought_signature": None,
                "signature_model": "gemini:gemini-2.5-flash",
            }]},
            {"role": "tool", "name": "query_data", "content": "3 rows"},
        ]
        original, self.HISTORY = self.HISTORY, history
        try:
            parts = await self.parts_sent_to("gemini:gemini-2.5-flash")
        finally:
            self.HISTORY = original

        calls = [p for p in parts if getattr(p, "function_call", None)]
        self.assertEqual(len(calls), 1, "the call is still replayed as a call")
        self.assertTrue([p for p in parts if getattr(p, "function_response", None)])

    async def test_a_history_with_no_tool_calls_is_unaffected(self):
        from unittest.mock import MagicMock

        from app.services.llm import LLMService

        service = LLMService()
        service._initialized = True
        service.gemini_client = MagicMock()
        captured = []
        answer = json.dumps({"answer": "hi", "visualizations": [],
                             "key_insights": [], "sources_used": []})

        class Models:
            def generate_content(self, model, contents, config):
                captured.append(contents)
                part = MagicMock(text=answer, function_call=None, thought_signature=None)
                candidate = MagicMock()
                candidate.content.parts = [part]
                return MagicMock(candidates=[candidate])

        service.gemini_client.models = Models()
        await service._chat_gemini(
            [{"role": "user", "content": "hello"}], "sys", None, 0.1, True, "gemini:any"
        )
        texts = [p.text for c in captured[0] for p in c.parts if getattr(p, "text", None)]
        self.assertEqual(texts, ["hello"])


class IdentifiersAreNotMeasures(unittest.TestCase):
    """
    `00067` is a label, not a number.

    Converting it lost the padding — the answer listed customer 67 — and the column then
    joined the correlation matrix, where the model read a 0.06 coefficient as customer
    number moving with revenue.
    """

    def setUp(self):
        self.svc = ingestor()

    def test_zero_padded_values_stay_text(self):
        df = pd.DataFrame({"customer_id": ["00067", "00115", "00035"]})
        out = self.svc._smart_numeric_conversion(df.copy())
        self.assertFalse(pd.api.types.is_numeric_dtype(out["customer_id"]))
        self.assertEqual(out["customer_id"].tolist(), ["00067", "00115", "00035"])

    def test_columns_named_after_a_configured_token_stay_text(self):
        """The tokens are configuration, so the test reads them rather than repeating them."""
        for token in settings.FORCE_STRING_COLUMNS:
            for name in (token, f"customer_{token}", f"{token}_value"):
                with self.subTest(column=name):
                    df = pd.DataFrame({name: ["1", "2", "3"]})
                    out = self.svc._smart_numeric_conversion(df.copy())
                    self.assertFalse(pd.api.types.is_numeric_dtype(out[name]))

    def test_a_token_appearing_inside_a_word_does_not_match(self):
        """Substring matching would force "notes" to text for the sake of "no"."""
        df = pd.DataFrame({"identity_score": ["1", "2"], "codex_rank": ["3", "4"]})
        out = self.svc._smart_numeric_conversion(df.copy())
        self.assertTrue(pd.api.types.is_numeric_dtype(out["identity_score"]))
        self.assertTrue(pd.api.types.is_numeric_dtype(out["codex_rank"]))

    def test_real_measures_still_convert(self):
        for name in ("revenue", "units", "margin_pct", "total_paid"):
            with self.subTest(column=name):
                df = pd.DataFrame({name: ["10", "20", "30"]})
                out = self.svc._smart_numeric_conversion(df.copy())
                self.assertTrue(pd.api.types.is_numeric_dtype(out[name]))

    def test_a_leading_zero_decimal_is_still_a_number(self):
        """0.5 starts with a zero but is not an identifier."""
        df = pd.DataFrame({"rate": ["0.5", "0.25", "0.75"]})
        out = self.svc._smart_numeric_conversion(df.copy())
        self.assertTrue(pd.api.types.is_numeric_dtype(out["rate"]))


class TheProfileNamesTheDataTraps(unittest.TestCase):
    def test_casing_variants_are_reported_with_the_real_count(self):
        from app.services.tools import _create_data_profile

        df = pd.DataFrame({
            "region": ["North", "north", "NORTH", " North", "South", "south"],
            "revenue": [1, 2, 3, 4, 5, 6],
        })
        profile = _create_data_profile(df, "sales")
        self.assertIn("spellings of 2 distinct values", profile)
        self.assertIn("Normalise before grouping", profile)

    def test_a_clean_column_gets_no_warning(self):
        from app.services.tools import _create_data_profile

        df = pd.DataFrame({"region": ["North", "South"], "revenue": [1, 2]})
        self.assertNotIn("spellings of", _create_data_profile(df, "sales"))

    def test_a_free_text_column_does_not_trigger_it(self):
        """High-cardinality prose would warn on every near-miss and drown the profile."""
        from app.services.tools import _casing_variants

        notes = pd.Series([f"note number {i}" for i in range(500)])
        self.assertEqual(_casing_variants(notes), "")


class TheRequestAdaptsToTheModel(unittest.IsolatedAsyncioTestCase):
    """
    Every OpenAI call sent temperature=0.1 and a json_schema response format.

    Reasoning-grade models reject any temperature but their default, so selecting one
    would have 400'd on the first call and failed over — the recommended models would have
    been the ones that did not work. Which parameter a given model refuses changes with
    every release, so the refusal is read from the provider's own complaint and remembered
    rather than kept in a compatibility matrix.
    """

    class Refuses(Exception):
        """A 400 shaped like the provider's, naming the parameter it will not accept."""

        def __init__(self, parameter):
            super().__init__(
                f"Error code: 400 - Unsupported value: '{parameter}' does not support "
                "0.1 with this model. Only the default (1) is supported."
            )
            self.status_code = 400

    def service(self, refuse=None, fail_times=1):
        from unittest.mock import MagicMock

        from app.services.llm import LLMService

        service = LLMService()
        service._initialized = True
        attempts = []

        async def create(**kwargs):
            attempts.append(dict(kwargs))
            if refuse and refuse in kwargs and len(attempts) <= fail_times:
                raise self.Refuses(refuse)
            return MagicMock(choices=[MagicMock(message=MagicMock(content="ok", tool_calls=None))])

        service.openai_client = MagicMock()
        service.openai_client.chat.completions.create = create
        return service, attempts

    async def test_a_refused_parameter_is_dropped_and_the_call_retried(self):
        service, attempts = self.service(refuse="temperature")
        await service._openai_create({"model": "gpt-5.2", "messages": [], "temperature": 0.1})

        self.assertEqual(len(attempts), 2)
        self.assertIn("temperature", attempts[0])
        self.assertNotIn("temperature", attempts[1])

    async def test_the_refusal_is_remembered_for_that_model(self):
        """Paying the round trip once is tolerable; paying it every call is not."""
        service, attempts = self.service(refuse="temperature")
        for _ in range(3):
            await service._openai_create(
                {"model": "gpt-5.2", "messages": [], "temperature": 0.1}
            )

        # First call refuses and retries; the two after it skip the parameter outright
        self.assertEqual(len(attempts), 4)
        self.assertNotIn("temperature", attempts[-1])

    async def test_what_one_model_refuses_does_not_affect_another(self):
        service, attempts = self.service(refuse="temperature")
        await service._openai_create({"model": "gpt-5.2", "messages": [], "temperature": 0.1})
        attempts.clear()
        await service._openai_create({"model": "gpt-4o", "messages": [], "temperature": 0.1})
        self.assertIn("temperature", attempts[0])

    async def test_any_droppable_parameter_is_handled(self):
        for parameter in ("response_format", "max_completion_tokens", "tool_choice"):
            with self.subTest(parameter=parameter):
                service, attempts = self.service(refuse=parameter)
                await service._openai_create(
                    {"model": "some-model", "messages": [], parameter: "x"}
                )
                self.assertNotIn(parameter, attempts[-1])

    async def test_a_400_naming_nothing_we_sent_is_not_swallowed(self):
        """Dropping a parameter at random would turn a real error into a silent one."""
        service, _ = self.service()

        async def always_400(**kwargs):
            raise self.Refuses("something_we_never_send")

        service.openai_client.chat.completions.create = always_400
        with self.assertRaises(self.Refuses):
            await service._openai_create({"model": "m", "messages": [], "temperature": 0.1})

    async def test_a_non_400_is_raised_untouched(self):
        """A rate limit must reach the retry and failover logic, not be reinterpreted."""
        from app.services.llm import LLMService

        service, _ = self.service()

        class RateLimited(Exception):
            status_code = 429

        async def limited(**kwargs):
            raise RateLimited("slow down")

        service.openai_client.chat.completions.create = limited
        with self.assertRaises(RateLimited):
            await service._openai_create({"model": "m", "messages": [], "temperature": 0.1})


class NoChunkIsUnbounded(unittest.TestCase):
    """
    A dataframe rendered to text has no blank lines, so splitting on paragraphs alone
    left it whole: a 6,000-row CSV became a single 258,000-character chunk. Every later
    retrieval then embedded a quarter of a megabyte and cross-encoded the same, which took
    one RAG query from 30 seconds to 8m22s before it was killed at ten minutes.
    """

    def setUp(self):
        self.svc = ingestor()

    def test_text_without_blank_lines_is_still_split(self):
        chunks = self.svc._semantic_chunking("x" * 250_000, "t", "d", 1)
        self.assertGreater(len(chunks), 100)
        for chunk in chunks:
            self.assertLessEqual(len(chunk["content"]), settings.CHUNK_SIZE * 1.5)

    def test_splitting_keeps_the_content(self):
        chunks = self.svc._semantic_chunking("x" * 250_000, "t", "d", 1)
        self.assertGreaterEqual(sum(len(c["content"]) for c in chunks), 250_000)

    def test_ordinary_prose_is_unaffected(self):
        prose = "\n\n".join(f"Paragraph {i}. Some ordinary text here." for i in range(60))
        chunks = self.svc._semantic_chunking(prose, "t", "d", 1)
        self.assertLess(len(chunks), 10)
        for chunk in chunks:
            self.assertLessEqual(len(chunk["content"]), settings.CHUNK_SIZE * 1.5)

    def test_a_large_table_indexes_a_preview_not_a_transcript(self):
        """
        The rows are queryable through the tools, and far more accurately. Indexing all
        of them buys nothing and costs an embedding over the whole file.
        """
        frame = pd.DataFrame({
            "id": range(6000),
            "amount": np.arange(6000, dtype=float),
            "branch": ["A", "B", "C", "D", "E", "F"] * 1000,
        })
        text = self.svc._table_as_text(frame, "transactions")

        self.assertLess(len(text), 6000, "a preview, not the whole frame")
        self.assertIn("6,000 rows", text)
        self.assertIn("branch", text, "the column names must stay searchable")
        self.assertIn("further rows", text, "and it must say what it left out")

    def test_a_small_table_is_shown_whole(self):
        frame = pd.DataFrame({"region": ["N", "S"], "revenue": [1.0, 2.0]})
        text = self.svc._table_as_text(frame, "sales")
        self.assertNotIn("further rows", text)
        self.assertIn("2 rows", text)
